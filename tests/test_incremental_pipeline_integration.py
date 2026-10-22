from __future__ import annotations

from pathlib import Path

import pymupdf as fitz
from docx import Document as WordDocument
from PIL import Image
from sqlalchemy.orm import Session, sessionmaker

from corporate_rag_ingestion.chunking.chunker import StructuralChunker
from corporate_rag_ingestion.chunking.overrides import ChunkingOverrides
from corporate_rag_ingestion.chunking.settings import ChunkingSettings
from corporate_rag_ingestion.chunking.storage import ChunkStorage
from corporate_rag_ingestion.config import Settings
from corporate_rag_ingestion.incremental.ocr import IncrementalOCRProcessor
from corporate_rag_ingestion.incremental.preparer import IncrementalDocumentPreparer
from corporate_rag_ingestion.incremental.service import IncrementalUpdater
from corporate_rag_ingestion.ocr.cache import OCRCache
from corporate_rag_ingestion.ocr.settings import MonkeyOCRSettings
from corporate_rag_ingestion.pipeline import IngestionPipeline
from corporate_rag_ingestion.retrieval.indexer import HybridIndexer
from corporate_rag_ingestion.retrieval.models import EmbeddingRuntime, SparseEmbedding
from corporate_rag_ingestion.retrieval.settings import RetrievalSettings
from corporate_rag_ingestion.retrieval.store import QdrantStore
from corporate_rag_ingestion.storage import NormalizedStorage


class TinyDense:
    def __init__(self, settings: RetrievalSettings) -> None:
        self.settings = settings

    @property
    def runtime(self) -> EmbeddingRuntime:
        return EmbeddingRuntime(
            model_id="test",
            model_revision="test",
            dimension=2,
            device="cpu",
            dtype="float32",
            settings_hash=self.settings.dense_settings_hash,
        )

    def embed_documents(self, texts):  # type: ignore[no-untyped-def]
        return [[1.0, 0.0] for _ in texts]

    def token_counts(self, texts):  # type: ignore[no-untyped-def]
        return [len(text.split()) for text in texts]


class TinySparse:
    def embed_documents(self, texts):  # type: ignore[no-untyped-def]
        return [SparseEmbedding(indices=[1], values=[1.0]) for _ in texts]


def save_docx(path: Path, text: str) -> None:
    document = WordDocument()
    document.add_heading("Регламент", level=1)
    document.add_paragraph(text)
    document.save(path)


def save_pdf(path: Path, text: str) -> None:
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_text((72, 100), text, fontsize=14)
    pdf.save(path)
    pdf.close()


def make_real_updater(
    tmp_path: Path,
    session_factory: sessionmaker[Session],
) -> tuple[IncrementalUpdater, QdrantStore, Path]:
    knowledge_base = tmp_path / "knowledge-base"
    knowledge_base.mkdir()
    storage_root = tmp_path / "storage"
    settings = Settings(
        _env_file=None,
        knowledge_base_path=knowledge_base,
        storage_path=storage_root,
        database_url=f"sqlite:///{(tmp_path / 'unused.db').as_posix()}",
        media_overrides_path=tmp_path / "no-media-overrides.json",
        chunk_overrides_path=tmp_path / "no-chunk-overrides.json",
        monkeyocr_enabled=False,
    )
    retrieval = RetrievalSettings(
        _env_file=None,
        qdrant_path=tmp_path / "qdrant",
        collection_name="incremental-real-pipeline",
        dense_dimension=2,
        dense_model_id="test",
        dense_model_revision="test",
        sparse_model_id="test",
        sparse_model_revision="test",
        sparse_cache_path=tmp_path / "sparse",
        model_cache_path=tmp_path / "models",
    )
    store = QdrantStore(retrieval)
    normalized = NormalizedStorage(storage_root)
    chunk_settings = ChunkingSettings()
    ocr_settings = MonkeyOCRSettings()
    preparer = IncrementalDocumentPreparer(
        ingestion=IngestionPipeline(settings, session_factory),
        normalized_storage=normalized,
        chunk_storage=ChunkStorage(storage_root / "chunks"),
        chunker=StructuralChunker(chunk_settings, ChunkingOverrides()),
        chunking_settings=chunk_settings,
        ocr=IncrementalOCRProcessor(
            storage=normalized,
            knowledge_base_root=knowledge_base,
            cache=OCRCache(storage_root / "ocr"),
            settings=ocr_settings,
            project_root=tmp_path,
        ),
    )
    updater = IncrementalUpdater(
        knowledge_base_root=knowledge_base,
        storage_root=storage_root,
        session_factory=session_factory,
        preparer=preparer,
        indexer=HybridIndexer(retrieval, store, TinyDense(retrieval), TinySparse()),  # type: ignore[arg-type]
        store=store,
    )
    return updater, store, knowledge_base


def test_real_docx_and_png_incremental_lifecycle(
    tmp_path: Path,
    session_factory: sessionmaker[Session],
) -> None:
    updater, store, root = make_real_updater(tmp_path, session_factory)
    try:
        docx = root / "policy.docx"
        save_docx(docx, "Старое нормативное значение — 10 дней.")
        added = updater.run()
        assert added.new == 1 and added.normalized == 1
        assert any("10 дней" in payload["text"] for payload in store.payloads().values())

        save_docx(docx, "Новое нормативное значение — 7 дней.")
        changed = updater.run()
        texts = [str(payload["text"]) for payload in store.payloads().values()]
        assert changed.changed == 1
        assert any("7 дней" in text for text in texts)
        assert all("10 дней" not in text for text in texts)

        image = root / "Сканирование ладони.png"
        Image.new("RGB", (80, 50), "white").save(image)
        image_added = updater.run()
        assert image_added.new == 1
        assert any(
            payload["relative_path"] == "Сканирование ладони.png"
            for payload in store.payloads().values()
        )

        pdf = root / "new-policy.pdf"
        save_pdf(
            pdf,
            "New PDF production policy 2026. " * 12,
        )
        pdf_added = updater.run()
        assert pdf_added.new == 1 and pdf_added.normalized == 1
        assert any(
            payload["relative_path"] == "new-policy.pdf"
            for payload in store.payloads().values()
        )

        docx.unlink()
        removed = updater.run()
        assert removed.deleted == 1
        assert all(
            payload["relative_path"] != "policy.docx"
            for payload in store.payloads().values()
        )

        repeated = updater.run()
        assert repeated.new == repeated.changed == repeated.deleted == 0
        assert repeated.unchanged == 2
        assert repeated.embeddings_calculated == 0
    finally:
        store.close()
