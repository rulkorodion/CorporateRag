from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from ..schemas import NormalizedElement
from .base import AdapterContext, FormatAdapter


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return str(value)


class ExcelAdapter(FormatAdapter):
    def normalize(self, context: AdapterContext):  # type: ignore[no-untyped-def]
        workbook = load_workbook(context.source_path, data_only=False, read_only=False)
        elements: list[NormalizedElement] = []
        sheets_metadata: list[dict[str, Any]] = []
        try:
            for order, sheet in enumerate(workbook.worksheets):
                cells: list[dict[str, Any]] = []
                min_row = sheet.max_row
                max_row = 0
                min_col = sheet.max_column
                max_col = 0
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell, MergedCell) or cell.value is None:
                            continue
                        min_row = min(min_row, cell.row)
                        max_row = max(max_row, cell.row)
                        min_col = min(min_col, cell.column)
                        max_col = max(max_col, cell.column)
                        cells.append(
                            {
                                "coordinate": cell.coordinate,
                                "row": cell.row,
                                "column": cell.column,
                                "value": json_value(cell.value),
                                "data_type": cell.data_type,
                            }
                        )
                used_range = None
                if cells:
                    from openpyxl.utils import get_column_letter

                    used_range = (
                        f"{get_column_letter(min_col)}{min_row}:"
                        f"{get_column_letter(max_col)}{max_row}"
                    )
                merged_cells = [str(item) for item in sheet.merged_cells.ranges]
                merged_cell_values = []
                for merged_range in sheet.merged_cells.ranges:
                    value = sheet.cell(merged_range.min_row, merged_range.min_col).value
                    if value is None:
                        continue
                    merged_cell_values.append(
                        {
                            "range": str(merged_range),
                            "min_row": merged_range.min_row,
                            "max_row": merged_range.max_row,
                            "min_column": merged_range.min_col,
                            "max_column": merged_range.max_col,
                            "value": json_value(value),
                        }
                    )
                tables = [
                    {"name": table.name, "display_name": table.displayName, "ref": table.ref}
                    for table in sheet.tables.values()
                ]
                sheet_data = {
                    "sheet_name": sheet.title,
                    "used_range": used_range,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "cells": cells,
                    "merged_cells": merged_cells,
                    "merged_cell_values": merged_cell_values,
                    "tables": tables,
                }
                sheets_metadata.append(
                    {
                        "sheet_name": sheet.title,
                        "used_range": used_range,
                        "cell_count": len(cells),
                        "merged_cells": merged_cells,
                        "tables": tables,
                    }
                )
                elements.append(
                    NormalizedElement(
                        element_id=f"{context.version_id}:sheet:{order}",
                        type="table",
                        order=order,
                        metadata=sheet_data,
                    )
                )
        finally:
            workbook.close()
        return context.normalized_document(
            metadata={
                **context.metadata,
                "workbook": context.source_path.name,
                "sheets": sheets_metadata,
            },
            elements=elements,
        )
