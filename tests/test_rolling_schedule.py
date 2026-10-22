from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from datetime import date
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from corporate_rag_ingestion.api.models import AnswerRequest
from corporate_rag_ingestion.generation.models import LLMGenerationResult
from corporate_rag_ingestion.rolling_schedule.normalizer import (
    EXCEL_ERRORS,
    ScheduleValidationError,
    extract_dimensions,
    normalize_workbook,
)
from corporate_rag_ingestion.rolling_schedule.answer import RollingScheduleAnswerService
from corporate_rag_ingestion.rolling_schedule.intent import (
    merge_intent_analysis,
    parse_intent_response,
)
from corporate_rag_ingestion.rolling_schedule.search import (
    is_schedule_candidate,
    is_schedule_intent,
    parse_schedule_query,
)
from corporate_rag_ingestion.rolling_schedule.settings import (
    EXPECTED_SHEETS,
    RollingScheduleSettings,
)
from corporate_rag_ingestion.rolling_schedule.source import KNOWN_GIDS, download_file
from corporate_rag_ingestion.rolling_schedule.store import RollingScheduleStore
from corporate_rag_ingestion.rolling_schedule.sync import RollingScheduleSyncService


def _settings(root: Path) -> RollingScheduleSettings:
    return RollingScheduleSettings(
        enabled=True,
        root=root,
        export_url="https://example.invalid/schedule.xlsx",
        edit_url="https://docs.google.com/spreadsheets/d/test/edit",
        spreadsheet_id="test",
        update_interval_hours=24,
        stale_after_hours=36,
        timeout_seconds=5,
        retries=3,
        max_download_bytes=10_000_000,
    )


def _create_fixture(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    working_sheets = [
        name
        for name in EXPECTED_SHEETS
        if name not in {"89", "250", "ПИ", "Лист ДМЦ ГК", "Лист ДМЦ ХК"}
    ]
    for index, name in enumerate(EXPECTED_SHEETS):
        sheet = workbook.create_sheet(name)
        if name in {"89", "250", "ПИ"}:
            for column in range(2, 5):
                sheet.cell(1, column, "#REF!")
            continue
        if name.startswith("Лист ДМЦ"):
            for offset in range(1, 7):
                row = 6 + offset
                if offset != 2:
                    sheet.cell(row, 1, date(2026, 10, 15))
                sheet.cell(row, 2, f"Лист 0,{offset}х1250х2500 ст3")
                sheet.cell(row, 3, 10 + offset)
                sheet.cell(row, 5, "ст3" if offset != 2 else "#VALUE!")
            continue
        # Five rows on 24 ordinary working sheets keep the fixture above the
        # production-file sanity threshold without depending on blank templates.
        for offset in range(1, 6):
            row = 5 + offset
            sheet.cell(row, 2, date(2026, 10, 15))
            if name == "114" and offset == 1:
                product = "Труба 57х3,5х6000 ст3"
            elif name == "114" and offset == 2:
                product = "Труба 57х3х6000 ст3"
            else:
                product = f"Труба {70 + index}х{offset},0х6000 ст3"
            sheet.cell(row, 3, product)
            sheet.cell(row, 4, 20 + offset)
            sheet.cell(row, 14, "#VALUE!" if offset == 3 else "ст3")
            sheet.cell(row, 16, date(2026, 10, 15))
        sheet.cell(20, 3, "Ремонт стана")
    assert len(working_sheets) == 24
    workbook.save(path)
    workbook.close()


@pytest.fixture()
def schedule_snapshot(tmp_path: Path) -> tuple[RollingScheduleSettings, Path]:
    source = tmp_path / "schedule.xlsx"
    _create_fixture(source)
    settings = _settings(tmp_path / "rolling")
    report = RollingScheduleSyncService(settings).sync(
        force=True, source_file=source, gid_map=dict(KNOWN_GIDS)
    )
    assert report.status == "updated"
    return settings, source


def test_normalizes_all_expected_sheets_and_formula_errors(tmp_path: Path) -> None:
    source = tmp_path / "schedule.xlsx"
    _create_fixture(source)
    result = normalize_workbook(
        source,
        settings=_settings(tmp_path / "rolling"),
        gid_map=dict(KNOWN_GIDS),
    )

    assert result.workbook_sheet_count == 29
    assert {item.sheet_name for item in result.sheets} == set(EXPECTED_SHEETS)
    diagnostics = {item.sheet_name: item for item in result.sheets}
    assert diagnostics["89"].status == "unavailable"
    assert diagnostics["250"].status == "unavailable"
    assert diagnostics["ПИ"].status == "unavailable"
    assert diagnostics["Лист ДМЦ ХК"].formula_errors["#VALUE!"] == 1
    dmc = next(
        item
        for item in result.records
        if item.sheet_name == "Лист ДМЦ ХК" and item.source_row == 8
    )
    assert dmc.product_name == "Лист 0,2х1250х2500 ст3"
    assert dmc.thickness == "0.2"
    assert dmc.width == "1250"
    assert dmc.length == "2500"
    assert "helper_formula_error" in dmc.uncertainties
    assert dmc.steel == "Ст3"
    assert dmc.volume == "12"


def test_exact_size_date_stand_and_source_link(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    store = RollingScheduleStore(settings.database_path, settings.status_path)

    query = parse_schedule_query(
        "Когда на стане 114 планируется прокат трубы 57×3,5 15.10.2026?",
        today=date(2026, 10, 14),
    )
    records = store.search(query)
    assert records
    assert {item.size_key for item in records} == {"57x3.5"}
    assert {item.sheet_name for item in records} == {"114"}
    assert all(item.start_date == date(2026, 10, 15) for item in records)
    assert all("gid=1050618602" in item.source_url for item in records)
    assert all("range=B6:P6" in item.source_url for item in records)

    close_but_different = store.search(
        parse_schedule_query(
            "Когда на стане 114 планируется прокат трубы 57x3?",
            today=date(2026, 10, 14),
        )
    )
    assert {item.size_key for item in close_but_different} == {"57x3"}
    assert {item.source_row for item in close_but_different} == {7}


def test_volume_query_keeps_trustworthy_undated_rows(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    store = RollingScheduleStore(settings.database_path, settings.status_path)

    with_date = parse_schedule_query(
        "Какой объём запланирован по листу 0,2×1250×2500 15.10.2026?",
        today=date(2026, 10, 14),
    )
    assert with_date.include_undated is False

    without_date = parse_schedule_query(
        "Какой объём запланирован по листу 0,2×1250×2500?",
        today=date(2026, 10, 16),
    )
    assert without_date.include_undated is True
    records = store.search(without_date)
    assert records
    assert {item.volume for item in records} == {"12"}


def test_sha_noop_and_failed_update_preserve_last_good(
    schedule_snapshot: tuple[RollingScheduleSettings, Path], tmp_path: Path
) -> None:
    settings, source = schedule_snapshot
    service = RollingScheduleSyncService(settings)
    before = hashlib.sha256(settings.database_path.read_bytes()).hexdigest()

    unchanged = service.sync(
        force=True, source_file=source, gid_map=dict(KNOWN_GIDS)
    )
    assert unchanged.status == "unchanged"
    assert unchanged.changed is False

    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_bytes(b"not an xlsx")
    failed = service.sync(force=True, source_file=corrupt)
    assert failed.status == "failed"
    assert hashlib.sha256(settings.database_path.read_bytes()).hexdigest() == before
    assert RollingScheduleStore(settings.database_path, settings.status_path).available


def test_download_retries_and_rejects_non_xlsx(tmp_path: Path) -> None:
    calls = 0

    def handler(_request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        if calls < 3:
            return httpx.Response(503, content=b"temporary")
        return httpx.Response(200, content=b"not a zip", headers={"content-type": "text/html"})

    client = httpx.Client(transport=httpx.MockTransport(handler))
    with pytest.raises(Exception, match="not an XLSX"):
        download_file(
            "https://example.invalid/file",
            tmp_path / "download.xlsx",
            timeout_seconds=5,
            retries=3,
            max_bytes=1000,
            client=client,
        )
    assert calls == 3
    client.close()


def test_incomplete_workbook_is_rejected(tmp_path: Path) -> None:
    incomplete = tmp_path / "incomplete.xlsx"
    workbook = Workbook()
    workbook.active.title = "159"
    workbook.save(incomplete)
    workbook.close()

    with pytest.raises(ScheduleValidationError, match="missing expected sheets"):
        normalize_workbook(
            incomplete,
            settings=_settings(tmp_path / "rolling"),
            gid_map=dict(KNOWN_GIDS),
        )


def test_intent_and_dimension_parsing_do_not_hijack_regular_rag() -> None:
    assert extract_dimensions("труба 57х3,5") == ("57", "3.5")
    assert extract_dimensions("труба 57×3") == ("57", "3")
    assert is_schedule_intent("Какие трубы запланированы на завтра?")
    assert is_schedule_intent("Что планируется на стане 114?")
    assert is_schedule_intent("План прокатов для стана 114")
    assert is_schedule_intent("Какие ремонты или простои запланированы на стане 114?")
    assert is_schedule_intent("Что будут выпускать на 114 стане на следующей неделе?")
    assert is_schedule_intent("Когда будут катать трубу 76х4?")
    assert is_schedule_intent("Какая загрузка Корейца на этой неделе?")
    assert is_schedule_intent("Что сейчас стоит в плане производства?")
    assert not is_schedule_intent("Как оформить командировку?")
    assert not is_schedule_intent("Какой регламент проката документов?")
    assert not is_schedule_intent("Как согласовать производственный заказ?")
    assert is_schedule_candidate("Что будем производить завтра?")
    assert is_schedule_candidate("Что стоит в плане на 15 октября?")
    assert not is_schedule_candidate("Как оформить командировку?")


def test_natural_production_question_extracts_stand_and_next_week() -> None:
    query = parse_schedule_query(
        "Что будут выпускать на 114 стане на следующей неделе?",
        today=date(2026, 10, 14),
    )
    assert query.is_schedule_intent
    assert query.sheet_names == ("114",)
    assert query.date_from == date(2026, 10, 19)
    assert query.date_to == date(2026, 10, 25)
    assert query.products_only is True

    inflected = parse_schedule_query(
        "Какая загрузка Корейца на этой неделе?",
        today=date(2026, 10, 14),
    )
    assert inflected.sheet_names == ("Кореец",)
    assert inflected.date_from == date(2026, 10, 14)
    assert inflected.date_to == date(2026, 10, 18)

    month = parse_schedule_query(
        "Что стоит в плане производства на октябрь?",
        today=date(2026, 10, 14),
    )
    assert month.date_from == date(2026, 10, 1)
    assert month.date_to == date(2026, 10, 31)

    period = parse_schedule_query(
        "Что будут выпускать на стане 114 с 15.10 по 20.10?",
        today=date(2026, 10, 14),
    )
    assert period.date_from == date(2026, 10, 15)
    assert period.date_to == date(2026, 10, 20)

    named_period = parse_schedule_query(
        "Покажи график прокатов с 15 по 17 октября",
        today=date(2026, 10, 14),
    )
    assert named_period.date_from == date(2026, 10, 15)
    assert named_period.date_to == date(2026, 10, 17)
    assert named_period.response_mode == "full"

    plan = parse_schedule_query(
        "План производства на 15 октября",
        today=date(2026, 10, 14),
    )
    assert plan.query_type == "production_plan"
    assert plan.response_mode == "full"

    stand_plan = parse_schedule_query(
        "График проката Ситко",
        today=date(2026, 10, 14),
    )
    assert stand_plan.query_type == "stand_schedule"
    assert stand_plan.response_mode == "full"

    text_summary = parse_schedule_query(
        "Кратко опиши план на 15 октября без таблицы",
        today=date(2026, 10, 14),
    )
    assert text_summary.response_mode == "brief"

    nearest = parse_schedule_query(
        "Какие позиции будут прокатывать в ближайшие три дня?",
        today=date(2026, 10, 14),
    )
    assert nearest.date_from == date(2026, 10, 14)
    assert nearest.date_to == date(2026, 10, 16)
    assert nearest.products_only is True

    stand = parse_schedule_query(
        "Какой график у стана 159 с 15 по 20 октябя?",
        today=date(2026, 10, 14),
    )
    assert stand.sheet_names == ("159",)
    assert stand.date_from == date(2026, 10, 15)
    assert stand.date_to == date(2026, 10, 20)
    assert stand.products_only is True

    open_from = parse_schedule_query(
        "Сколько производственных позиций запланировано с 15 октября 2026 года "
        "до конца доступного графика?",
        today=date(2026, 10, 15),
    )
    assert open_from.query_type == "count"
    assert open_from.date_from == date(2026, 10, 15)
    assert open_from.date_to is None
    assert open_from.period_mode == "open_from"

    open_after = parse_schedule_query(
        "Сколько позиций в плане проката после 20 октя 2026 года?",
        today=date(2026, 10, 15),
    )
    assert open_after.date_from == date(2026, 10, 21)
    assert open_after.date_to is None
    assert open_after.period_mode == "open_after"

    open_until = parse_schedule_query(
        "Сколько позиций в плане проката до 20 октя 2026 года?",
        today=date(2026, 10, 15),
    )
    assert open_until.date_from is None
    assert open_until.date_to == date(2026, 10, 20)
    assert open_until.period_mode == "open_until"

    all_available = parse_schedule_query(
        "Сколько всего производственных позиций во всех графиках, "
        "включая прошедшие даты?",
        today=date(2026, 10, 15),
    )
    assert all_available.date_from is None and all_available.date_to is None
    assert all_available.period_mode == "all_available"
    assert all_available.response_mode == "brief"


def test_service_event_query_does_not_return_products(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    query = parse_schedule_query(
        "Какие ремонты или простои запланированы на стане 114?",
        today=date(2026, 10, 14),
    )
    assert query.record_kinds == ("service",)
    assert query.include_undated is True
    records = RollingScheduleStore(
        settings.database_path, settings.status_path
    ).search(query)
    assert records
    assert {item.record_kind for item in records} == {"service"}
    assert {item.product_name for item in records} == {"Ремонт стана"}
    assert {item.start_date for item in records} == {None}


def test_grounded_fallback_keeps_only_public_schedule_fields(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    response = RollingScheduleAnswerService(settings).try_answer(
        "request-1",
        "Когда на стане 114 планируется прокат трубы 57×3,5 15.10.2026?",
        today=date(2026, 10, 14),
    )
    assert response is not None
    assert response.answer_kind == "rolling_schedule"
    assert response.model == "structured-rolling-schedule-fallback"
    assert "57х3,5х6000" in response.answer
    assert "15.10.2026" in response.answer
    assert response.sources == []
    assert response.source_attachments == []
    assert "http" not in response.answer
    assert "Плановый объём" not in response.answer
    assert "Источник" not in response.answer

    assert (
        RollingScheduleAnswerService(settings).try_answer(
            "request-2", "Как оформить командировку?"
        )
        is None
    )


def test_schedule_plan_is_grounded_for_llm_and_keeps_backend_table(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    service = RollingScheduleAnswerService(settings)
    plan = service.prepare(
        "request-llm",
        "Какие позиции есть в графике прокатов?",
        today=date(2026, 10, 14),
    )
    assert plan is not None
    assert plan.llm_request is not None
    assert '"date_from": "2026-10-15"' in plan.llm_request.user_prompt
    assert '"maximum_examples_to_list": 6' in plan.llm_request.user_prompt
    assert '"position_counts_by_stand"' in plan.llm_request.user_prompt
    assert '"source_url"' not in plan.llm_request.user_prompt
    assert len(plan.supplied_records) <= 10

    final = service.finish(
        plan,
        LLMGenerationResult(
            content="На стане 114 запланирован выпуск нескольких позиций.",
            model="test-model",
            done_reason="stop",
        ),
        llm_queue_wait_seconds=0.2,
        llm_queue_position=1,
        generation_seconds=0.5,
    )
    assert final.model == "test-model"
    assert "Источники:" not in final.answer
    assert "| Стан / площадка | Дата | Номенклатура |" in final.answer
    assert "gid=" not in final.answer
    assert "Плановый объём" not in final.answer
    assert "Последняя синхронизация" in final.answer
    assert final.latency.llm_queue_wait_seconds == 0.2
    assert final.telemetry.llm_queue_position == 1
    assert final.sources == []
    assert final.source_attachments == []


def test_plan_request_attaches_backend_table_by_default(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    plan = RollingScheduleAnswerService(settings).prepare(
        "request-full-date",
        "График прокатов на 15 октября 2026",
        today=date(2026, 10, 14),
    )
    assert plan is not None
    assert plan.llm_request is not None
    assert len(plan.records) > 20
    assert "Позиции плана — страница 1" in plan.appendix
    assert "Показаны строки **1–25" in plan.appendix
    assert "повторите запрос" not in plan.appendix
    assert plan.fallback.schedule_pagination is not None
    assert plan.fallback.schedule_pagination.has_next is True
    assert "Показаны 4 примера" not in plan.fallback.answer


def test_explicit_full_result_is_paginated_without_silent_truncation(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    plan = RollingScheduleAnswerService(settings).prepare(
        "request-full-date",
        "Покажи весь график прокатов на 15 октября 2026, страница 1",
        today=date(2026, 10, 14),
    )
    assert plan is not None
    assert plan.query.response_mode == "full"
    assert plan.llm_request is not None
    assert "Позиции плана — страница 1" in plan.appendix
    assert "Показаны строки **1–25" in plan.appendix
    assert plan.fallback.schedule_pagination is not None
    assert plan.fallback.schedule_pagination.page == 1


def test_named_stand_includes_undated_positions_in_compact_answer_context(
    tmp_path: Path,
) -> None:
    source = tmp_path / "schedule.xlsx"
    _create_fixture(source)
    workbook = load_workbook(source)
    sheet = workbook["Ситко"]
    for row in range(6, 11):
        sheet.cell(row, 2).value = None
        sheet.cell(row, 16).value = None
    workbook.save(source)
    workbook.close()

    settings = _settings(tmp_path / "rolling")
    report = RollingScheduleSyncService(settings).sync(
        force=True, source_file=source, gid_map=dict(KNOWN_GIDS)
    )
    assert report.status == "updated"
    plan = RollingScheduleAnswerService(settings).prepare(
        "request-full-stand",
        "План производства Ситко",
        today=date(2026, 10, 14),
    )
    assert plan is not None
    assert plan.llm_request is not None
    assert len(plan.records) == 5
    assert "дата не подтверждена" in plan.fallback.answer
    assert len(plan.supplied_records) == 5
    assert "| Стан / площадка | Дата | Номенклатура |" in plan.appendix
    assert plan.fallback.sources == []
    assert plan.fallback.source_attachments == []


def test_undated_service_events_are_not_called_confirmed_plan(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    plan = RollingScheduleAnswerService(settings).prepare(
        "request-service",
        "Какие ремонты или простои запланированы на стане 114?",
        today=date(2026, 10, 14),
    )
    assert plan is not None
    assert plan.llm_request is not None
    assert "даты этих событий в источнике не подтверждены" in plan.fallback.answer.lower()


def test_service_recognition_includes_cleaning_and_roll_change(tmp_path: Path) -> None:
    source = tmp_path / "services.xlsx"
    _create_fixture(source)
    workbook = load_workbook(source)
    sheet = workbook["114"]
    sheet.cell(21, 3, "ППР/ чистка")
    sheet.cell(22, 3, "перевалки 11ч")
    sheet.cell(23, 3, "1/6 перевалки")
    workbook.save(source)
    workbook.close()

    result = normalize_workbook(
        source,
        settings=_settings(tmp_path / "rolling"),
        gid_map=dict(KNOWN_GIDS),
    )
    services = [
        item.product_name
        for item in result.records
        if item.sheet_name == "114" and item.record_kind == "service"
    ]
    assert {"ППР/ чистка", "перевалки 11ч", "1/6 перевалки"} <= set(services)


def test_llm_intent_is_strictly_validated_and_deterministic_facts_win() -> None:
    base = parse_schedule_query(
        "Какой объём трубы 57×3,5 запланирован на стане 114?",
        today=date(2026, 10, 14),
    )
    analysis = parse_intent_response(
        '{"is_schedule_question":true,"query_type":"product_search",'
        '"stands":["114","250"],"product_kind":"round_tube",'
        '"dimensions":["57","4"],"steel":null,"period_text":null,'
        '"response_mode":"brief","needs_clarification":false,'
        '"clarification_question":null}'
    )
    assert analysis is not None
    merged = merge_intent_analysis(base, analysis)
    assert merged is not None
    assert merged.sheet_names == ("114",)
    assert merged.dimensions == ("57", "3.5")
    assert merged.query_type == "product_search"
    assert merged.products_only is True

    assert parse_intent_response('{"query_type":"volume"}') is None
    assert parse_intent_response("not-json") is None


def test_mixed_question_requests_clarification() -> None:
    base = parse_schedule_query(
        "Что можем производить и что запланировано на завтра?",
        today=date(2026, 10, 14),
    )
    analysis = parse_intent_response(
        '{"is_schedule_question":true,"query_type":"mixed","stands":[],'
        '"product_kind":null,"dimensions":[],"steel":null,'
        '"period_text":"на завтра","response_mode":"brief",'
        '"needs_clarification":true,"clarification_question":"Уточните тип данных."}'
    )
    merged = merge_intent_analysis(base, analysis)
    assert merged is not None
    assert merged.needs_clarification is True
    assert merged.query_type == "mixed"


@pytest.mark.parametrize(
    ("question", "query_type"),
    [
        ("График производства по стану Ситко", "stand_schedule"),
        ("График прокатов", "production_plan"),
        ("С какой по какую дату план производства?", "date_bounds"),
        ("Какие даты присутствуют в плане производства? Дай полный перечень", "date_list"),
    ],
)
def test_clear_schedule_questions_cannot_be_overridden_by_llm_clarification(
    question: str,
    query_type: str,
) -> None:
    base = parse_schedule_query(question, today=date(2026, 10, 15))
    analysis = parse_intent_response(
        '{"is_schedule_question":true,"query_type":"mixed","stands":[],'
        '"product_kind":null,"dimensions":[],"steel":null,"period_text":null,'
        '"response_mode":"brief","needs_clarification":true,'
        '"clarification_question":"Уточните технологические возможности или план."}'
    )
    merged = merge_intent_analysis(base, analysis)
    assert merged is not None
    assert merged.is_schedule_intent is True
    assert merged.query_type == query_type
    assert merged.needs_clarification is False


def test_date_bounds_and_date_list_are_exact_backend_answers(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    service = RollingScheduleAnswerService(settings)
    bounds = service.prepare(
        "date-bounds",
        "С какой по какую дату план производства?",
        today=date(2026, 10, 15),
    )
    dates = service.prepare(
        "date-list",
        "Какие даты присутствуют в плане производства? Дай полный перечень",
        today=date(2026, 10, 15),
    )
    assert bounds is not None and dates is not None
    assert bounds.llm_request is None
    assert dates.llm_request is None
    assert "15.10.2026" in bounds.fallback.answer
    assert "уникальных дат" in dates.fallback.answer
    assert "15.10.2026" in dates.fallback.answer


def test_schedule_pagination_is_snapshot_bound_and_resets_after_update(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    service = RollingScheduleAnswerService(settings)
    first = service.prepare(
        "first-page",
        "График прокатов на 15 октября 2026",
        today=date(2026, 10, 14),
    )
    assert first is not None
    pagination = first.fallback.schedule_pagination
    assert pagination is not None
    assert len(pagination.snapshot_version) == 64
    assert pagination.selection.date_from == date(2026, 10, 15)

    second = service.prepare(
        "second-page",
        "График прокатов на 15 октября 2026",
        today=date(2026, 10, 14),
        query=replace(first.query, page=2),
        expected_snapshot_version=pagination.snapshot_version,
        pagination_only=True,
    )
    assert second is not None
    assert second.llm_request is None
    assert second.fallback.schedule_pagination is not None
    assert second.fallback.schedule_pagination.page == 2

    changed = service.prepare(
        "stale-page",
        "График прокатов на 15 октября 2026",
        today=date(2026, 10, 14),
        query=replace(first.query, page=2),
        expected_snapshot_version="0" * 64,
        pagination_only=True,
    )
    assert changed is not None
    assert changed.fallback.schedule_pagination is not None
    assert changed.fallback.schedule_pagination.page == 1
    assert changed.fallback.schedule_pagination.snapshot_changed is True
    assert "График обновился" in changed.fallback.answer


def test_schedule_page_request_requires_complete_selection_state() -> None:
    with pytest.raises(ValidationError):
        AnswerRequest(
            question="График прокатов Ситко",
            schedule_page=2,
            schedule_snapshot_version="a" * 64,
        )

    request = AnswerRequest(
        question="График прокатов Ситко",
        schedule_page=2,
        schedule_snapshot_version="a" * 64,
        schedule_selection={
            "sheet_names": ["Ситко"],
            "date_from": "2026-10-15",
            "period_mode": "default_future",
            "include_undated": True,
            "products_only": True,
            "query_type": "stand_schedule",
            "response_mode": "full",
        },
    )
    assert request.schedule_selection is not None
    assert request.schedule_selection.date_from == date(2026, 10, 15)


def test_volume_question_does_not_expose_quantities_or_sources(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    service = RollingScheduleAnswerService(settings)
    plan = service.prepare(
        "request-volume",
        "Какой объём трубы 57×3,5 запланирован?",
        today=date(2026, 10, 14),
    )
    assert plan is not None and plan.llm_request is not None
    assert '"type": "product_search"' in plan.llm_request.user_prompt
    assert '"volume_aggregates"' not in plan.llm_request.user_prompt
    assert '"volume"' not in plan.llm_request.user_prompt
    assert '"unit"' not in plan.llm_request.user_prompt
    assert '"source_row"' not in plan.llm_request.user_prompt
    assert plan.fallback.sources == []
    assert plan.fallback.source_attachments == []
    assert "Плановый объём" not in plan.fallback.answer
    assert "Источник" not in plan.fallback.answer

    final = service.finish(
        plan,
        LLMGenerationResult(
            content="Запланированный объём — 120 тонн.",
            model="test-model",
            done_reason="stop",
        ),
        llm_queue_wait_seconds=0,
        llm_queue_position=1,
        generation_seconds=0.1,
    )
    assert final.model == "structured-rolling-schedule-fallback"
    assert "120 тонн" not in final.answer


def test_unverified_llm_fact_falls_back_to_structured_answer(
    schedule_snapshot: tuple[RollingScheduleSettings, Path],
) -> None:
    settings, _ = schedule_snapshot
    service = RollingScheduleAnswerService(settings)
    plan = service.prepare(
        "request-grounding",
        "Когда на стане 114 планируется труба 57×3,5?",
        today=date(2026, 10, 14),
    )
    assert plan is not None and plan.llm_request is not None
    final = service.finish(
        plan,
        LLMGenerationResult(
            content="Запланирована труба 58×3,5.",
            model="test-model",
            done_reason="stop",
        ),
        llm_queue_wait_seconds=0,
        llm_queue_position=1,
        generation_seconds=0.1,
    )
    assert final.model == "structured-rolling-schedule-fallback"
    assert "58×3,5" not in final.answer
    assert final.sources == []
    assert final.source_attachments == []


def test_real_verified_workbook_when_available(tmp_path: Path) -> None:
    real = Path(r"C:\Users\Rodion\Downloads\prokat_test.xlsx")
    if not real.is_file():
        pytest.skip("local verified workbook is not available")
    result = normalize_workbook(
        real,
        settings=_settings(tmp_path / "rolling"),
        gid_map=dict(KNOWN_GIDS),
    )
    diagnostics = {item.sheet_name: item for item in result.sheets}
    assert result.workbook_sheet_count == 29
    assert diagnostics["Лист ДМЦ ГК"].product_count == 21
    assert diagnostics["Лист ДМЦ ХК"].product_count == 45
    assert diagnostics["114"].service_count == 9
    assert diagnostics["89"].status == "unavailable"
    assert diagnostics["250"].status == "unavailable"
    assert diagnostics["ПИ"].status == "unavailable"
    assert any(
        item.sheet_name == "Лист ДМЦ ХК"
        and item.source_row == 34
        and item.thickness == "0.66"
        and "helper_formula_error" in item.uncertainties
        for item in result.records
    )
    assert not any(
        value in EXCEL_ERRORS
        for item in result.records
        for value in (item.volume, item.steel)
    )
    service_names = {
        item.product_name
        for item in result.records
        if item.sheet_name == "114" and item.record_kind == "service"
    }
    assert {"ППР/ чистка", "перевалки 11ч", "1/6 перевалки"} <= service_names


def test_real_workbook_user_questions_use_validated_structured_filters(
    tmp_path: Path,
) -> None:
    real = Path(r"C:\Users\Rodion\Downloads\prokat_test.xlsx")
    if not real.is_file():
        pytest.skip("local verified workbook is not available")
    settings = _settings(tmp_path / "real-rolling")
    report = RollingScheduleSyncService(settings).sync(
        force=True, source_file=real, gid_map=dict(KNOWN_GIDS)
    )
    assert report.status == "updated"
    store = RollingScheduleStore(settings.database_path, settings.status_path)
    fixed_today = date(2026, 10, 14)

    questions = [
        "Что будем производить завтра?",
        "Что стоит в плане на 15 октября?",
        "Покажи график прокатов с 15 по 17 октября.",
        "Что планируется на этой неделе?",
        "Какие позиции будут прокатывать в ближайшие три дня?",
        "Какой график у стана 159 с 15 по 20 октября?",
        "Что запланировано на Ижевск ХК завтра?",
        "Где и когда будут делать профиль 60×40×2?",
        "Когда запланирован лист 0,8×1250×2500?",
        "Какой объём трубы 57×3,5 запланирован?",
        "Какие ремонты и ППР запланированы на этой неделе?",
    ]
    parsed = [parse_schedule_query(item, today=fixed_today) for item in questions]
    assert all(item.is_schedule_intent for item in parsed)
    assert all(store.search(item) for item in parsed)
    assert parsed[0].date_from == parsed[0].date_to == date(2026, 10, 15)
    assert (parsed[2].date_from, parsed[2].date_to) == (
        date(2026, 10, 15),
        date(2026, 10, 17),
    )
    assert (parsed[4].date_from, parsed[4].date_to) == (
        date(2026, 10, 14),
        date(2026, 10, 16),
    )
    assert parsed[5].sheet_names == ("159",)
    assert parsed[6].sheet_names == ("Ижевск ХК",)
    assert parsed[7].dimensions == ("60", "40", "2")
    assert parsed[8].dimensions == ("0.8", "1250", "2500")
    assert parsed[9].dimensions == ("57", "3.5")
    assert parsed[10].record_kinds == ("service",)
    assert all(
        item.record_kind == "service" for item in store.search(parsed[10])
    )
    assert not is_schedule_candidate("Как оформить перемещение?")

    service = RollingScheduleAnswerService(settings)
    plan_on_date = service.prepare(
        "plan-on-date",
        "План производства на 15 октября",
        today=date(2026, 10, 14),
    )
    full_plan_on_date = service.prepare(
        "full-plan-on-date",
        "Полный план производства на 15 октября",
        today=date(2026, 10, 14),
    )
    second_page_on_date = service.prepare(
        "second-page-on-date",
        "План производства на 15 октября, страница 2",
        today=date(2026, 10, 14),
    )
    tomorrow_plan = service.prepare(
        "tomorrow-plan",
        "Что будем производить завтра?",
        today=date(2026, 10, 14),
    )
    sitko_plan = service.prepare(
        "sitko-plan",
        "График прокатов Ситко",
        today=date(2026, 10, 14),
    )
    count_on_date = service.prepare(
        "count-on-date",
        "Сколько позиций на 15 октября?",
        today=date(2026, 10, 14),
    )
    brief_on_date = service.prepare(
        "brief-on-date",
        "Кратко опиши план на 15 октября без таблицы",
        today=date(2026, 10, 14),
    )

    assert plan_on_date is not None and full_plan_on_date is not None
    assert second_page_on_date is not None
    assert tomorrow_plan is not None and sitko_plan is not None
    assert count_on_date is not None and brief_on_date is not None
    assert len(plan_on_date.records) == 39
    assert len(full_plan_on_date.records) == 39
    assert len(tomorrow_plan.records) == 39
    assert plan_on_date.query.response_mode == "full"
    assert full_plan_on_date.query.response_mode == "full"
    assert tomorrow_plan.query.response_mode == "full"
    assert "страница 1 из 2" in plan_on_date.appendix
    assert "Показаны строки **1–25 из 39**" in plan_on_date.appendix
    assert "страница 2 из 2" in second_page_on_date.appendix
    assert "Показаны строки **26–39 из 39**" in second_page_on_date.appendix
    sitko_total = len(sitko_plan.records)
    sitko_pages = (sitko_total + 24) // 25
    assert sitko_total > 50
    assert f"страница 1 из {sitko_pages}" in sitko_plan.appendix
    assert f"Показаны строки **1–25 из {sitko_total}**" in sitko_plan.appendix
    assert len(count_on_date.records) == 39
    assert count_on_date.query.query_type == "count"
    assert count_on_date.appendix == ""
    assert "39" in count_on_date.fallback.answer
    assert brief_on_date.query.response_mode == "brief"
    assert brief_on_date.appendix == ""

    open_count = service.prepare(
        "open-count",
        "Сколько производственных позиций запланировано с 15 октября 2026 года "
        "до конца доступного графика?",
        today=date(2026, 10, 15),
    )
    assert open_count is not None
    assert open_count.query.date_to is None
    assert len(open_count.records) == 1067
    assert open_count.llm_request is None
    assert "1067" in open_count.fallback.answer
    assert "04.12.2026" in open_count.fallback.answer
    assert "| Дата |" not in open_count.fallback.answer

    same_selection = service.prepare(
        "open-count-rephrased",
        "Каково количество позиций в плане, начиная с 15 октября 2026 года?",
        today=date(2026, 10, 15),
    )
    assert same_selection is not None
    assert len(same_selection.records) == len(open_count.records)

    all_count = service.prepare(
        "all-count",
        "Сколько всего производственных позиций во всех графиках, "
        "включая прошедшие даты?",
        today=date(2026, 10, 15),
    )
    assert all_count is not None
    assert len(all_count.records) == 1830
    assert "1500" in all_count.fallback.answer
    assert "330" in all_count.fallback.answer
    assert "1500 + 330 = 1830" in all_count.fallback.answer
    assert "| Дата |" not in all_count.fallback.answer
    with_volume_and_unit = sum(
        item.volume is not None and item.unit is not None
        for item in all_count.records
    )
    with_volume_without_unit = sum(
        item.volume is not None and item.unit is None
        for item in all_count.records
    )
    without_volume = sum(item.volume is None for item in all_count.records)
    assert (with_volume_and_unit, with_volume_without_unit, without_volume) == (
        40,
        1533,
        257,
    )
    assert with_volume_and_unit + with_volume_without_unit + without_volume == 1830

    number_only = service.prepare(
        "number-only",
        "Сколько позиций в плане проката? Только цифры, без таблицы.",
        today=date(2026, 10, 15),
    )
    assert number_only is not None
    assert number_only.fallback.answer == "1067"
    assert number_only.appendix == ""

    tubes = service.prepare(
        "tubes-on-date",
        "Какие трубы запланированы на 15 октя?",
        today=date(2026, 10, 15),
    )
    assert tubes is not None and tubes.records
    assert tubes.llm_request is not None
    assert "| Стан / площадка | Дата | Номенклатура |" in tubes.appendix
    assert {item.record_kind for item in tubes.records} == {"product"}
    assert all(item.start_date <= date(2026, 10, 15) for item in tubes.records if item.start_date)
    assert all(
        (item.end_date or item.start_date) >= date(2026, 10, 15)
        for item in tubes.records
        if item.start_date
    )

    payload = json.loads(tubes.llm_request.user_prompt.split("\n\n", 1)[1])
    assert "volume_coverage" not in payload["matched_record_breakdown"]
    assert "volume_aggregates" not in payload
    assert all(
        set(item) == {"kind", "stand", "date_from", "date_to", "nomenclature"}
        for item in payload["examples"]
    )
