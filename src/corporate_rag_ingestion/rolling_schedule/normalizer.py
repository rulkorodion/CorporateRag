from __future__ import annotations

import hashlib
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from .models import NormalizationResult, ScheduleRecord, SheetDiagnostic
from .settings import RollingScheduleSettings


EXCEL_ERRORS = {
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#N/A",
    "#DIV/0!",
    "#NUM!",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
}
NORMALIZER_VERSION = "rolling-schedule-v2"
DIMENSION_RE = re.compile(
    r"(?<!\d)(\d{1,5}(?:[.,]\d+)?)\s*[xх×]\s*"
    r"(\d{1,5}(?:[.,]\d+)?)"
    r"(?:\s*[xх×]\s*(\d{1,5}(?:[.,]\d+)?))?"
    r"(?:\s*[xх×]\s*(\d{1,5}(?:[.,]\d+)?))?",
    re.IGNORECASE,
)
SERVICE_RE = re.compile(
    r"\b(выходн\w*|ремонт\w*|ппр|переналад\w*|простой\w*|"
    r"перевалк\w*|чистк\w*|настройк\w*|корректировк\w*|"
    r"профилактик\w*|останов\w*)\b",
    re.IGNORECASE,
)
PRODUCT_RE = re.compile(r"\b(труб\w*|профил\w*|лист\w*)\b", re.IGNORECASE)
STEEL_RE = re.compile(
    r"\b(09г2с|08пс|08кп|ст\s*\d{1,2}(?:пс|сп|кп)?|s\d{3,4})\b",
    re.IGNORECASE,
)
RUSSIAN_MONTHS = {
    "янв": 1,
    "фев": 2,
    "мар": 3,
    "апр": 4,
    "май": 5,
    "мая": 5,
    "июн": 6,
    "июл": 7,
    "авг": 8,
    "сен": 9,
    "окт": 10,
    "ноя": 11,
    "дек": 12,
}


class ScheduleValidationError(RuntimeError):
    pass


def is_excel_error(value: Any) -> bool:
    return isinstance(value, str) and value.strip().upper() in EXCEL_ERRORS


def normalize_number(value: Any) -> str | None:
    if value is None or isinstance(value, bool) or is_excel_error(value):
        return None
    if isinstance(value, Decimal):
        number = value
    elif isinstance(value, (int, float)):
        number = Decimal(str(value))
    elif isinstance(value, str):
        candidate = value.strip().replace(" ", "").replace(",", ".")
        if not candidate:
            return None
        try:
            number = Decimal(candidate)
        except InvalidOperation:
            return None
    else:
        return None
    if not number.is_finite():
        return None
    rendered = format(number.normalize(), "f")
    if "." in rendered:
        rendered = rendered.rstrip("0").rstrip(".")
    return "0" if rendered in {"-0", ""} else rendered


def extract_dimensions(text: str) -> tuple[str, ...]:
    match = DIMENSION_RE.search(text)
    if not match:
        return ()
    return tuple(
        value
        for value in (normalize_number(item) for item in match.groups())
        if value is not None
    )


def product_kind(text: str) -> str | None:
    lowered = text.casefold()
    if "лист" in lowered:
        return "sheet"
    if "проф" in lowered or re.search(r"\bпр\b", lowered):
        return "profile"
    if "труб" in lowered:
        return "round_tube"
    return "product" if extract_dimensions(text) else None


def size_dimensions(kind: str | None, dimensions: tuple[str, ...]) -> tuple[str, ...]:
    if kind == "round_tube":
        return dimensions[:2]
    if kind in {"profile", "sheet"}:
        return dimensions[:3]
    return dimensions


def dimension_fields(
    kind: str | None, dimensions: tuple[str, ...]
) -> dict[str, str | None]:
    values: dict[str, str | None] = {
        "diameter": None,
        "width": None,
        "height": None,
        "thickness": None,
        "length": None,
    }
    if kind == "round_tube":
        if len(dimensions) > 0:
            values["diameter"] = dimensions[0]
        if len(dimensions) > 1:
            values["thickness"] = dimensions[1]
        if len(dimensions) > 2:
            values["length"] = dimensions[2]
    elif kind == "profile":
        if len(dimensions) > 0:
            values["width"] = dimensions[0]
        if len(dimensions) > 1:
            values["height"] = dimensions[1]
        if len(dimensions) > 2:
            values["thickness"] = dimensions[2]
        if len(dimensions) > 3:
            values["length"] = dimensions[3]
    elif kind == "sheet":
        if len(dimensions) > 0:
            values["thickness"] = dimensions[0]
        if len(dimensions) > 1:
            values["width"] = dimensions[1]
        if len(dimensions) > 2:
            values["length"] = dimensions[2]
    return values


def parse_source_date(value: Any, *, default_year: int) -> tuple[date | None, bool]:
    if value is None or is_excel_error(value):
        return None, False
    if isinstance(value, datetime):
        return value.date(), False
    if isinstance(value, date):
        return value, False
    if not isinstance(value, str):
        return None, False
    text = value.strip().casefold().replace("г.", "").replace("г", "")
    if not text:
        return None, False
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date(), False
        except ValueError:
            pass
    numeric = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})", text)
    if numeric:
        try:
            return date(default_year, int(numeric.group(2)), int(numeric.group(1))), True
        except ValueError:
            return None, False
    russian = re.fullmatch(r"(\d{1,2})\s*[.]?\s*([а-яё]{3,})[.]?", text)
    if russian:
        month_token = russian.group(2)[:3]
        month = RUSSIAN_MONTHS.get(month_token)
        if month:
            try:
                return date(default_year, month, int(russian.group(1))), True
            except ValueError:
                return None, False
    return None, False


def _string(value: Any) -> str | None:
    if value is None or is_excel_error(value):
        return None
    text = " ".join(str(value).split())
    return text or None


def _raw_string(value: Any) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def _merged_value(sheet: Any, row: int, column: int) -> Any:
    value = sheet.cell(row, column).value
    if value is not None and not isinstance(sheet.cell(row, column), MergedCell):
        return value
    for merged in sheet.merged_cells.ranges:
        if (
            merged.min_row <= row <= merged.max_row
            and merged.min_col <= column <= merged.max_col
        ):
            return sheet.cell(merged.min_row, merged.min_col).value
    return value


def _workbook_default_year(workbook: Any) -> int:
    years: Counter[int] = Counter()
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_col=1, max_col=min(sheet.max_column, 16)):
            for cell in row:
                if isinstance(cell.value, (date, datetime)):
                    years[cell.value.year] += 1
    return years.most_common(1)[0][0] if years else datetime.now().year


def _source_url(edit_url: str, gid: str | None, row: int, start_col: str, end_col: str) -> str:
    if gid is None:
        return edit_url
    return f"{edit_url}#gid={quote(gid)}&range={start_col}{row}:{end_col}{row}"


def _record_id(spreadsheet_id: str, sheet_name: str, row: int, name: str) -> str:
    identity = f"{spreadsheet_id}\0{sheet_name}\0{row}\0{name}".encode("utf-8")
    return hashlib.sha256(identity).hexdigest()[:32]


def _steel(text: str, helper: Any) -> str | None:
    helper_text = _string(helper)
    if helper_text and helper_text not in {"__", "-"}:
        return helper_text
    match = STEEL_RE.search(text)
    if not match:
        return None
    return re.sub(r"\s+", "", match.group(1)).upper().replace("СТ", "Ст")


def _sheet_layout(sheet_name: str) -> tuple[int, int, int, int, int, str, str]:
    if sheet_name.startswith("Лист ДМЦ"):
        return 2, 3, 1, 13, 5, "A", "M"
    return 3, 4, 2, 16, 14, "B", "P"


def _sheet_unit(sheet: Any, volume_col: int) -> str | None:
    for row in range(1, min(sheet.max_row, 12) + 1):
        value = _string(sheet.cell(row, volume_col).value)
        if value and re.search(r"\b(?:тн|тонн\w*|т)\b", value.casefold()):
            return "т"
    return None


def normalize_workbook(
    path: Path,
    *,
    settings: RollingScheduleSettings,
    gid_map: dict[str, str],
) -> NormalizationResult:
    if not path.is_file() or not zipfile.is_zipfile(path):
        raise ScheduleValidationError("downloaded file is not a valid XLSX ZIP archive")
    try:
        workbook = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:
        raise ScheduleValidationError(f"openpyxl could not read schedule XLSX: {exc}") from exc
    try:
        missing = [name for name in settings.expected_sheets if name not in workbook.sheetnames]
        if missing:
            raise ScheduleValidationError(
                "schedule XLSX is missing expected sheets: " + ", ".join(missing)
            )
        default_year = _workbook_default_year(workbook)
        all_records: list[ScheduleRecord] = []
        diagnostics: list[SheetDiagnostic] = []
        for sheet_name in settings.expected_sheets:
            sheet = workbook[sheet_name]
            gid = gid_map.get(sheet_name)
            name_col, volume_col, start_col, end_col, steel_col, range_start, range_end = (
                _sheet_layout(sheet_name)
            )
            unit = _sheet_unit(sheet, volume_col)
            error_counts: Counter[str] = Counter()
            for row in sheet.iter_rows():
                for cell in row:
                    if is_excel_error(cell.value):
                        error_counts[str(cell.value).upper()] += 1
            import_broken = any(
                str(sheet.cell(1, column).value).upper() == "#REF!"
                for column in range(2, min(4, sheet.max_column) + 1)
            )
            records: list[ScheduleRecord] = []
            for row in range(1, sheet.max_row + 1):
                raw_name = sheet.cell(row, name_col).value
                name = _string(raw_name)
                if not name or name in {"-", "__"}:
                    continue
                dimensions = extract_dimensions(name)
                kind = product_kind(name)
                is_product = bool(dimensions and (PRODUCT_RE.search(name) or not SERVICE_RE.search(name)))
                is_service = bool(SERVICE_RE.search(name)) and not is_product
                if not is_product and not is_service:
                    continue
                raw_start = _merged_value(sheet, row, start_col)
                raw_end = _merged_value(sheet, row, end_col)
                start_date, start_year_inferred = parse_source_date(
                    raw_start, default_year=default_year
                )
                end_date, end_year_inferred = parse_source_date(
                    raw_end, default_year=default_year
                )
                if end_date is not None and start_date is not None and end_date < start_date:
                    end_date = None
                raw_volume = sheet.cell(row, volume_col).value
                volume = normalize_number(raw_volume) if is_product else None
                # Rows with a product-like catalogue name but without either a
                # schedule date or a plan value are not production positions.
                if (
                    is_product
                    and start_date is None
                    and end_date is None
                    and volume is None
                ):
                    continue
                uncertainties: list[str] = []
                if is_product and start_date is None:
                    uncertainties.append("date_missing")
                if is_product and volume is None:
                    uncertainties.append(
                        "volume_formula_error" if is_excel_error(raw_volume) else "volume_missing"
                    )
                if start_year_inferred or end_year_inferred:
                    uncertainties.append("date_year_inferred_from_workbook")
                helper_errors = 0
                for column in range(5 if not sheet_name.startswith("Лист ДМЦ") else 5, min(sheet.max_column, 15) + 1):
                    if is_excel_error(sheet.cell(row, column).value):
                        helper_errors += 1
                if helper_errors:
                    uncertainties.append("helper_formula_error")
                helper_values = [
                    sheet.cell(row, column).value
                    for column in range(steel_col, min(sheet.max_column, 15) + 1)
                ]
                if is_product and not any(
                    value is not None and not is_excel_error(value)
                    for value in helper_values
                ):
                    uncertainties.append("helper_values_missing")
                if is_product and kind is not None and not dimensions:
                    uncertainties.append("dimensions_missing")
                size = size_dimensions(kind, dimensions)
                fields = dimension_fields(kind, dimensions)
                source_url = _source_url(
                    settings.edit_url, gid, row, range_start, range_end
                )
                records.append(
                    ScheduleRecord(
                        record_id=_record_id(
                            settings.spreadsheet_id, sheet_name, row, name
                        ),
                        sheet_name=sheet_name,
                        gid=gid,
                        source_row=row,
                        record_kind="product" if is_product else "service",
                        start_date=start_date,
                        end_date=end_date,
                        product_name=name,
                        product_kind=kind if is_product else "service",
                        dimensions=dimensions,
                        size_key="x".join(size) if size else None,
                        steel=_steel(name, sheet.cell(row, steel_col).value)
                        if is_product
                        else None,
                        volume=volume,
                        unit=unit,
                        # Raw fields deliberately retain Excel error literals.
                        # They are diagnostics, never normalized production facts.
                        raw_date=_raw_string(raw_start),
                        raw_volume=_raw_string(raw_volume),
                        uncertainties=tuple(dict.fromkeys(uncertainties)),
                        source_url=source_url,
                        **fields,
                    )
                )
            products = [item for item in records if item.record_kind == "product"]
            services = [item for item in records if item.record_kind == "service"]
            warnings: list[str] = []
            status = "ok"
            if import_broken:
                status = "unavailable"
                warnings.append("B1:D1 contain #REF!: imported source data is unavailable")
                records = []
                products = []
                services = []
            elif not products and error_counts:
                status = "unavailable"
                warnings.append(
                    "no production rows were recovered while formula errors are present"
                )
            elif any(item.start_date is None for item in products):
                status = "partial"
                warnings.append("some production rows have no confirmed date")
            if any("helper_formula_error" in item.uncertainties for item in products):
                if status == "ok":
                    status = "partial"
                warnings.append(
                    "helper formulas contain errors; characteristics were parsed from source names"
                )
            helper_issue_count = sum(
                "helper_formula_error" in item.uncertainties
                or "helper_values_missing" in item.uncertainties
                for item in products
            )
            if any("helper_values_missing" in item.uncertainties for item in products):
                if status == "ok":
                    status = "partial"
                warnings.append(
                    "some helper characteristics are missing; only unambiguous values from source names are used"
                )
            dated_count = sum(item.start_date is not None for item in records)
            diagnostics.append(
                SheetDiagnostic(
                    sheet_name=sheet_name,
                    gid=gid,
                    status=status,
                    product_count=len(products),
                    service_count=len(services),
                    dated_record_count=dated_count,
                    undated_record_count=len(records) - dated_count,
                    formula_error_count=sum(error_counts.values()),
                    helper_issue_count=helper_issue_count,
                    formula_errors=dict(error_counts),
                    warnings=tuple(dict.fromkeys(warnings)),
                )
            )
            all_records.extend(records)
        product_total = sum(item.record_kind == "product" for item in all_records)
        if product_total < 100:
            raise ScheduleValidationError(
                f"schedule XLSX contains too few recognizable production rows: {product_total}"
            )
        return NormalizationResult(
            records=tuple(all_records),
            sheets=tuple(diagnostics),
            workbook_sheet_count=len(workbook.sheetnames),
        )
    finally:
        workbook.close()
